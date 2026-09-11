"""
app/services/export_service.py
------------------------------
In-memory document export service for study notes:
- Generates publication-quality, student-friendly PDF documents using ReportLab.
- Generates cleanly structured Word documents (.docx) using python-docx.
Returns bytes directly in-memory (using io.BytesIO) without writing scratch files to disk.
Supports AdaptiveStudyNotes, TopicStudyNotes, StudyNotes, and arbitrary dictionary payloads.
"""

import io
import html
import csv
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional, Union

# ReportLab imports
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import (
    HRFlowable,
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.pdfgen import canvas

# python-docx imports
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.ai.schemas import AdaptiveStudyNotes, StudyNotes, TopicDetailBlock
from app.services.study_generator import (
    ConceptBlock,
    TopicStudyNotes,
    MCQItem,
    ShortAnswerItem,
    GlossaryItem,
    CompleteStudyPack,
)


# ==============================================================================
# Normalized Intermediate Representation
# ==============================================================================

class NormalizedSection:
    def __init__(
        self,
        title: str,
        overview: str,
        key_points: List[str],
        code_or_syntax: Optional[str] = None,
    ):
        self.title = title
        self.overview = overview
        self.key_points = key_points
        self.code_or_syntax = code_or_syntax


class NormalizedNotes:
    def __init__(
        self,
        title: str,
        summary: str,
        sections: List[NormalizedSection],
        takeaways: List[str],
        pitfalls: Optional[List[str]] = None,
    ):
        self.title = title
        self.summary = summary
        self.sections = sections
        self.takeaways = takeaways
        self.pitfalls = pitfalls or []


def _normalize_notes(notes_input: Any) -> NormalizedNotes:
    """Safely convert any notes payload (AdaptiveStudyNotes, TopicStudyNotes, StudyNotes, dict) into NormalizedNotes."""
    if notes_input is None:
        return NormalizedNotes(title="Study Notes", summary="", sections=[], takeaways=[])

    # 1. AdaptiveStudyNotes
    if isinstance(notes_input, AdaptiveStudyNotes):
        sections = [
            NormalizedSection(
                title=s.title or "Section",
                overview=s.overview or "",
                key_points=[p for p in (s.key_points or []) if p and str(p).strip()],
                code_or_syntax=(s.code_or_syntax.strip() if s.code_or_syntax and str(s.code_or_syntax).strip() else None),
            )
            for s in (notes_input.sections or [])
        ]
        return NormalizedNotes(
            title=notes_input.title or "Study Notes",
            summary=notes_input.executive_summary or "",
            sections=sections,
            takeaways=[t for t in (notes_input.actionable_takeaways or []) if t and str(t).strip()],
            pitfalls=[p for p in (notes_input.common_pitfalls or []) if p and str(p).strip()],
        )

    # 2. TopicStudyNotes
    if isinstance(notes_input, TopicStudyNotes):
        sections = [
            NormalizedSection(
                title=c.subtopic or "Concept",
                overview=c.explanation or "",
                key_points=[c.practical_example] if (c.practical_example and c.practical_example.strip()) else [],
                code_or_syntax=(c.syntax_or_formula.strip() if c.syntax_or_formula and c.syntax_or_formula.strip() else None),
            )
            for c in (notes_input.concepts or [])
        ]
        return NormalizedNotes(
            title=notes_input.topic_title or "Study Notes",
            summary=notes_input.quick_summary or "",
            sections=sections,
            takeaways=[t for t in (notes_input.key_takeaways or []) if t and str(t).strip()],
            pitfalls=[p for p in (notes_input.common_pitfalls or []) if p and str(p).strip()],
        )

    # 3. StudyNotes
    if isinstance(notes_input, StudyNotes):
        sections = [
            NormalizedSection(
                title=c.term or "Concept",
                overview=c.definition or "",
                key_points=[],
                code_or_syntax=(c.syntax_or_example.strip() if c.syntax_or_example and c.syntax_or_example.strip() else None),
            )
            for c in (notes_input.core_concepts or [])
        ]
        return NormalizedNotes(
            title=notes_input.topic_title or "Study Notes",
            summary=notes_input.executive_summary or "",
            sections=sections,
            takeaways=[t for t in (notes_input.high_yield_revision_points or []) if t and str(t).strip()],
            pitfalls=[],
        )

    # 4. Dictionary
    if isinstance(notes_input, dict):
        title = (
            notes_input.get("title")
            or notes_input.get("topic_title")
            or notes_input.get("topic")
            or "Study Notes"
        )
        summary = (
            notes_input.get("executive_summary")
            or notes_input.get("quick_summary")
            or notes_input.get("summary")
            or ""
        )

        sections: List[NormalizedSection] = []
        if "sections" in notes_input and isinstance(notes_input["sections"], list):
            for s in notes_input["sections"]:
                if isinstance(s, dict):
                    kp = s.get("key_points") or []
                    if not isinstance(kp, list):
                        kp = [str(kp)]
                    code = s.get("code_or_syntax") or s.get("syntax_or_formula") or s.get("syntax_or_example")
                    sections.append(
                        NormalizedSection(
                            title=s.get("title") or s.get("subtopic") or s.get("term") or "Section",
                            overview=s.get("overview") or s.get("explanation") or s.get("definition") or "",
                            key_points=[str(p) for p in kp if p],
                            code_or_syntax=str(code).strip() if code and str(code).strip() else None,
                        )
                    )
        elif "concepts" in notes_input and isinstance(notes_input["concepts"], list):
            for c in notes_input["concepts"]:
                if isinstance(c, dict):
                    kp = [c["practical_example"]] if c.get("practical_example") else []
                    code = c.get("syntax_or_formula")
                    sections.append(
                        NormalizedSection(
                            title=c.get("subtopic") or c.get("title") or "Concept",
                            overview=c.get("explanation") or c.get("overview") or "",
                            key_points=[str(p) for p in kp if p],
                            code_or_syntax=str(code).strip() if code and str(code).strip() else None,
                        )
                    )
        elif "core_concepts" in notes_input and isinstance(notes_input["core_concepts"], list):
            for c in notes_input["core_concepts"]:
                if isinstance(c, dict):
                    code = c.get("syntax_or_example")
                    sections.append(
                        NormalizedSection(
                            title=c.get("term") or c.get("concept") or "Concept",
                            overview=c.get("definition") or "",
                            key_points=[],
                            code_or_syntax=str(code).strip() if code and str(code).strip() else None,
                        )
                    )

        raw_takeaways = (
            notes_input.get("actionable_takeaways")
            or notes_input.get("key_takeaways")
            or notes_input.get("high_yield_revision_points")
            or []
        )
        if not isinstance(raw_takeaways, list):
            raw_takeaways = [str(raw_takeaways)]

        raw_pitfalls = notes_input.get("common_pitfalls") or []
        if not isinstance(raw_pitfalls, list):
            raw_pitfalls = [str(raw_pitfalls)]

        return NormalizedNotes(
            title=str(title),
            summary=str(summary),
            sections=sections,
            takeaways=[str(t) for t in raw_takeaways if t and str(t).strip()],
            pitfalls=[str(p) for p in raw_pitfalls if p and str(p).strip()],
        )

    # 5. Fallback for unexpected objects
    t = getattr(notes_input, "title", getattr(notes_input, "topic_title", "Study Notes"))
    s = getattr(notes_input, "executive_summary", getattr(notes_input, "quick_summary", ""))
    return NormalizedNotes(title=str(t), summary=str(s), sections=[], takeaways=[])


class NumberedCanvas(canvas.Canvas):
    """Two-pass canvas to dynamically compute and print total page numbers."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            super().showPage()
        super().save()

    def draw_page_number(self, page_count: int):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))
        # Footer text
        footer_text = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(letter[0] - 36, 25, footer_text)
        self.drawString(36, 25, "AI Study Assistant — High-Yield Academic Revision Guide")
        # Thin divider line
        self.setStrokeColor(colors.HexColor("#CBD5E1"))
        self.setLineWidth(0.5)
        self.line(36, 35, letter[0] - 36, 35)
        self.restoreState()


# ==============================================================================
# PDF Export Engine (ReportLab)
# ==============================================================================

def export_notes_to_pdf(notes_input: Union[AdaptiveStudyNotes, TopicStudyNotes, StudyNotes, Dict[str, Any]]) -> io.BytesIO:
    """
    Generate an in-memory academic PDF study guide from any study notes structure.
    Returns an io.BytesIO buffer seeked to position 0.
    """
    notes = _normalize_notes(notes_input)
    buffer = io.BytesIO()

    # 36pt (0.5 inch) margins for maximized readable area
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=45,
    )

    styles = getSampleStyleSheet()

    # Color Palette
    PRIMARY = colors.HexColor("#1E3A8A")     # Deep Navy
    SECONDARY = colors.HexColor("#2563EB")   # Royal Blue
    TEXT_DARK = colors.HexColor("#0F172A")   # Slate 900
    TEXT_MUTED = colors.HexColor("#475569")  # Slate 600
    BOX_BG = colors.HexColor("#F8FAFC")      # Slate 50
    BOX_BORDER = colors.HexColor("#CBD5E1")  # Slate 300
    ALERT_BG = colors.HexColor("#FEF2F2")    # Red 50
    ALERT_BORDER = colors.HexColor("#F87171")# Red 400
    ALERT_TEXT = colors.HexColor("#991B1B")  # Red 800

    # Custom Typography Styles
    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=PRIMARY,
        spaceAfter=4,
    )

    subtitle_style = ParagraphStyle(
        "DocSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        textColor=TEXT_MUTED,
        spaceAfter=14,
    )

    h1_style = ParagraphStyle(
        "SectionHeader",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=16,
        textColor=SECONDARY,
        spaceBefore=12,
        spaceAfter=6,
        keepWithNext=True,
    )

    h2_style = ParagraphStyle(
        "SubSectionHeader",
        parent=styles["Heading3"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        textColor=PRIMARY,
        spaceBefore=8,
        spaceAfter=4,
        keepWithNext=True,
    )

    body_style = ParagraphStyle(
        "Body",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=13.5,
        textColor=TEXT_DARK,
        spaceAfter=6,
    )

    bullet_style = ParagraphStyle(
        "BulletPoint",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        textColor=TEXT_DARK,
        leftIndent=14,
        firstLineIndent=-10,
        spaceAfter=4,
    )

    code_style = ParagraphStyle(
        "CodeText",
        parent=styles["Normal"],
        fontName="Courier",
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor("#1E293B"),
    )

    story = []

    # Title & Metadata Header
    safe_title = html.escape(notes.title)
    story.append(Paragraph(safe_title, title_style))
    created_str = datetime.now(timezone.utc).strftime("%B %d, %Y")
    story.append(Paragraph(f"Adaptive Study Guide  |  Generated on {created_str}", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=PRIMARY, spaceAfter=10))

    # 1. Executive Summary Callout Box
    if notes.summary and notes.summary.strip():
        story.append(Paragraph("1. Executive Summary", h1_style))
        summary_p = Paragraph(html.escape(notes.summary.strip()).replace("\n", "<br/>"), body_style)
        summary_table = Table([[summary_p]], colWidths=[letter[0] - 72])
        summary_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EFF6FF")),  # Blue 50
            ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#BFDBFE")),        # Blue 200
            ("TOPPADDING", (0, 0), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ("LEFTPADDING", (0, 0), (-1, -1), 10),
            ("RIGHTPADDING", (0, 0), (-1, -1), 10),
        ]))
        story.append(summary_table)
        story.append(Spacer(1, 10))

    # 2. Detailed Breakdown (Sections / Modules)
    if notes.sections:
        story.append(Paragraph("2. Detailed Breakdown & Architecture", h1_style))
        for idx, sec in enumerate(notes.sections, start=1):
            sec_elements = []
            safe_sec_title = html.escape(sec.title)
            sec_elements.append(Paragraph(f"2.{idx}&nbsp; {safe_sec_title}", h2_style))

            # Overview / Explanation
            if sec.overview and sec.overview.strip():
                safe_ov = html.escape(sec.overview.strip()).replace("\n", "<br/>")
                sec_elements.append(Paragraph(safe_ov, body_style))

            # Key Points (Concrete requirements, principles, deliverables)
            if sec.key_points:
                for kp in sec.key_points:
                    safe_kp = html.escape(str(kp).strip())
                    sec_elements.append(Paragraph(f"&bull;&nbsp;&nbsp;{safe_kp}", bullet_style))

            # Code / Syntax / Formula Box (ONLY rendered if explicitly present in source text!)
            if sec.code_or_syntax and sec.code_or_syntax.strip():
                sec_elements.append(Spacer(1, 3))
                code_content = html.escape(sec.code_or_syntax.strip()).replace("\n", "<br/>")
                code_p = Paragraph(f"<b>CODE / SYNTAX / SPECIFICATION:</b><br/>{code_content}", code_style)
                code_table = Table([[code_p]], colWidths=[letter[0] - 72])
                code_table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, -1), BOX_BG),
                    ("BOX", (0, 0), (-1, -1), 0.75, BOX_BORDER),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ("LEFTPADDING", (0, 0), (-1, -1), 8),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ]))
                sec_elements.append(code_table)

            sec_elements.append(Spacer(1, 8))
            story.append(KeepTogether(sec_elements))

    # 3. Actionable Takeaways & Implementation Checkpoints
    if notes.takeaways:
        story.append(Paragraph("3. Actionable Takeaways & Implementation Checkpoints", h1_style))
        for point in notes.takeaways:
            safe_point = html.escape(str(point).strip())
            story.append(Paragraph(f"&bull;&nbsp;&nbsp;<b>{safe_point}</b>", bullet_style))
        story.append(Spacer(1, 10))

    # 4. Common Pitfalls & Failure Modes (Rendered ONLY if non-empty!)
    if notes.pitfalls:
        pitfalls_elements = [
            Paragraph("4. Common Traps & Failure Modes", h1_style)
        ]
        for pitfall in notes.pitfalls:
            safe_pit = html.escape(str(pitfall).strip())
            pitfall_p = Paragraph(f"<b>&#9888;&nbsp; {safe_pit}</b>", ParagraphStyle(
                "PitfallText",
                parent=bullet_style,
                textColor=ALERT_TEXT,
            ))
            pitfalls_elements.append(pitfall_p)

        pitfalls_table = Table([[pitfalls_elements]], colWidths=[letter[0] - 72])
        pitfalls_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), ALERT_BG),
            ("BOX", (0, 0), (-1, -1), 1, ALERT_BORDER),
            ("TOPPADDING", (0, 0), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ("LEFTPADDING", (0, 0), (-1, -1), 10),
            ("RIGHTPADDING", (0, 0), (-1, -1), 10),
        ]))
        story.append(Spacer(1, 4))
        story.append(KeepTogether([pitfalls_table]))

    doc.build(story, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    return buffer


# ==============================================================================
# Word (.docx) Export Engine (python-docx)
# ==============================================================================

def _set_cell_shading(cell, color_hex: str):
    """Set background color of a Word table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tc_pr.append(shd)


def export_notes_to_docx(notes_input: Union[AdaptiveStudyNotes, TopicStudyNotes, StudyNotes, Dict[str, Any]]) -> io.BytesIO:
    """
    Generate an in-memory formatted Word (.docx) study guide from any study notes structure.
    Returns an io.BytesIO buffer seeked to position 0.
    """
    notes = _normalize_notes(notes_input)
    doc = Document()

    # Set page margins to 0.75 in
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # 1. Document Title
    title_p = doc.add_paragraph()
    title_run = title_p.add_run(notes.title)
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    title_p.paragraph_format.space_after = Pt(2)

    # Subtitle
    sub_p = doc.add_paragraph()
    created_str = datetime.now(timezone.utc).strftime("%B %d, %Y")
    sub_run = sub_p.add_run(f"Adaptive Study Guide  |  Generated on {created_str}")
    sub_run.font.name = "Calibri"
    sub_run.font.size = Pt(9.5)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    sub_p.paragraph_format.space_after = Pt(14)

    # 2. Executive Summary
    if notes.summary and notes.summary.strip():
        h1 = doc.add_heading("1. Executive Summary", level=1)
        h1.paragraph_format.space_before = Pt(12)
        h1.paragraph_format.space_after = Pt(4)

        # Render summary in shaded callout box
        table = doc.add_table(rows=1, cols=1)
        table.autofit = False
        table.columns[0].width = Inches(7.0)
        cell = table.cell(0, 0)
        _set_cell_shading(cell, "EFF6FF")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(notes.summary.strip())
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(6)

    # 3. Detailed Breakdown (Sections / Modules)
    if notes.sections:
        h1 = doc.add_heading("2. Detailed Breakdown & Architecture", level=1)
        h1.paragraph_format.space_before = Pt(12)
        h1.paragraph_format.space_after = Pt(6)

        for idx, sec in enumerate(notes.sections, start=1):
            h2 = doc.add_heading(f"2.{idx} {sec.title}", level=2)
            h2.paragraph_format.space_before = Pt(8)
            h2.paragraph_format.space_after = Pt(3)

            # Overview
            if sec.overview and sec.overview.strip():
                exp_p = doc.add_paragraph()
                r = exp_p.add_run(sec.overview.strip())
                r.font.name = "Calibri"
                r.font.size = Pt(10)
                exp_p.paragraph_format.space_after = Pt(4)

            # Key Points
            if sec.key_points:
                for kp in sec.key_points:
                    bp = doc.add_paragraph(style="List Bullet")
                    r = bp.add_run(str(kp).strip())
                    r.font.name = "Calibri"
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
                    bp.paragraph_format.space_after = Pt(2)

            # Code / Syntax / Formula Callout (ONLY rendered if explicitly present!)
            if sec.code_or_syntax and sec.code_or_syntax.strip():
                tbl = doc.add_table(rows=1, cols=1)
                tbl.autofit = False
                tbl.columns[0].width = Inches(7.0)
                cell = tbl.cell(0, 0)
                _set_cell_shading(cell, "F1F5F9")
                p = cell.paragraphs[0]
                hdr_run = p.add_run("CODE / SYNTAX / SPECIFICATION:\n")
                hdr_run.font.bold = True
                hdr_run.font.size = Pt(8.5)
                hdr_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
                code_run = p.add_run(sec.code_or_syntax.strip())
                code_run.font.name = "Consolas"
                code_run.font.size = Pt(9)
                code_run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
                p.paragraph_format.space_after = Pt(4)

            doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 4. Actionable Takeaways
    if notes.takeaways:
        h1 = doc.add_heading("3. Actionable Takeaways & Implementation Checkpoints", level=1)
        h1.paragraph_format.space_before = Pt(12)
        h1.paragraph_format.space_after = Pt(4)

        for point in notes.takeaways:
            bp = doc.add_paragraph(style="List Bullet")
            r = bp.add_run(str(point).strip())
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            bp.paragraph_format.space_after = Pt(3)

        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(6)

    # 5. Common Pitfalls (ONLY rendered if non-empty!)
    if notes.pitfalls:
        h1 = doc.add_heading("4. Common Traps & Failure Modes", level=1)
        h1.paragraph_format.space_before = Pt(12)
        h1.paragraph_format.space_after = Pt(4)

        table = doc.add_table(rows=1, cols=1)
        table.autofit = False
        table.columns[0].width = Inches(7.0)
        cell = table.cell(0, 0)
        _set_cell_shading(cell, "FEF2F2")

        for idx, pitfall in enumerate(notes.pitfalls):
            p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
            warn_icon = p.add_run("⚠ ")
            warn_icon.font.bold = True
            warn_icon.font.color.rgb = RGBColor(0xDC, 0x26, 0x26)
            r = p.add_run(str(pitfall).strip())
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x99, 0x1B, 0x1B)
            p.paragraph_format.space_after = Pt(3)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def export_mcqs_to_csv(mcqs: List[Any]) -> io.StringIO:
    """
    Export MCQs to a standard CSV format compatible with Anki, Quizlet, and flashcard tools.
    Columns: Question, Options, Correct Answer, Explanation
    Options format: A. ... | B. ... | C. ... | D. ...
    """
    output = io.StringIO()
    writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
    writer.writerow(["Question", "Options", "Correct Answer", "Explanation"])
    for item in mcqs:
        q = getattr(item, "question", None) or (item.get("question") if isinstance(item, dict) else "")
        opts = getattr(item, "options", None) or (item.get("options") if isinstance(item, dict) else [])
        ans = getattr(item, "correct_answer", None) or (item.get("correct_answer") if isinstance(item, dict) else "")
        expl = getattr(item, "explanation", None) or (item.get("explanation") if isinstance(item, dict) else "")

        # If correct_answer was not set but correct_index was
        if not ans and opts:
            idx = getattr(item, "correct_index", None) or (item.get("correct_index") if isinstance(item, dict) else None)
            if idx is not None and 0 <= idx < len(opts):
                ans = opts[idx]

        options_str = " | ".join([f"{chr(65+i)}. {opt}" for i, opt in enumerate(opts)])
        writer.writerow([q, options_str, ans, expl])

    output.seek(0)
    return output


def export_pack_to_pdf(pack: Union[CompleteStudyPack, Dict[str, Any]]) -> io.BytesIO:
    """
    Generate an in-memory academic PDF study pack from CompleteStudyPack using ReportLab.
    Returns an io.BytesIO buffer seeked to position 0.
    """
    if isinstance(pack, dict):
        pack = CompleteStudyPack.model_validate(pack)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=45,
    )

    styles = getSampleStyleSheet()

    PRIMARY = colors.HexColor("#1E3A8A")     # Deep Navy
    SECONDARY = colors.HexColor("#2563EB")   # Royal Blue
    TEXT_DARK = colors.HexColor("#0F172A")   # Slate 900
    TEXT_MUTED = colors.HexColor("#475569")  # Slate 600
    BOX_BG = colors.HexColor("#F8FAFC")      # Slate 50
    BOX_BORDER = colors.HexColor("#CBD5E1")  # Slate 300

    title_style = ParagraphStyle(
        "PackTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=PRIMARY,
        spaceAfter=4,
    )

    subtitle_style = ParagraphStyle(
        "PackSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        textColor=TEXT_MUTED,
        spaceAfter=12,
    )

    h1_style = ParagraphStyle(
        "PackSectionHeader",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=16,
        textColor=SECONDARY,
        spaceBefore=14,
        spaceAfter=6,
        keepWithNext=True,
    )

    h2_style = ParagraphStyle(
        "PackSubSectionHeader",
        parent=styles["Heading3"],
        fontName="Helvetica-Bold",
        fontSize=10.5,
        leading=13.5,
        textColor=PRIMARY,
        spaceBefore=8,
        spaceAfter=4,
        keepWithNext=True,
    )

    body_style = ParagraphStyle(
        "PackBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=13.5,
        textColor=TEXT_DARK,
        spaceAfter=6,
    )

    bullet_style = ParagraphStyle(
        "PackBullet",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        textColor=TEXT_DARK,
        leftIndent=14,
        firstLineIndent=-10,
        spaceAfter=4,
    )

    table_header_style = ParagraphStyle(
        "TableHeader",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9,
        leading=12,
        textColor=colors.whitesmoke,
    )

    table_cell_style = ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11.5,
        textColor=TEXT_DARK,
    )

    story = []

    # 1. Header: Title, Difficulty Badge, Generation Date
    safe_title = html.escape(pack.title)
    story.append(Paragraph(safe_title, title_style))
    created_str = datetime.now(timezone.utc).strftime("%B %d, %Y")
    diff_label = pack.difficulty.upper()
    story.append(Paragraph(
        f"<b>Complete Academic Study Pack</b> &nbsp;|&nbsp; "
        f"Difficulty: <b>{diff_label}</b> &nbsp;|&nbsp; Generated on {created_str}",
        subtitle_style,
    ))
    story.append(HRFlowable(width="100%", thickness=1.5, color=PRIMARY, spaceAfter=10))

    # 2. Suggested Study Order
    if pack.suggested_study_order:
        story.append(Paragraph("1. Suggested Study Progression & Roadmap", h1_style))
        for idx, step in enumerate(pack.suggested_study_order, start=1):
            safe_step = html.escape(str(step).strip())
            story.append(Paragraph(f"<b>Step {idx}:</b>&nbsp; {safe_step}", bullet_style))
        story.append(Spacer(1, 8))

    # 3. Concise Summary Notes
    if pack.concise_summary:
        story.append(Paragraph("2. Concise Subject Summary Notes", h1_style))
        safe_summary = html.escape(pack.concise_summary.strip()).replace("\n", "<br/>")
        summary_p = Paragraph(safe_summary, body_style)
        summary_table = Table([[summary_p]], colWidths=[letter[0] - 72])
        summary_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EFF6FF")),
            ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#BFDBFE")),
            ("TOPPADDING", (0, 0), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ("LEFTPADDING", (0, 0), (-1, -1), 10),
            ("RIGHTPADDING", (0, 0), (-1, -1), 10),
        ]))
        story.append(summary_table)
        story.append(Spacer(1, 10))

    # 4. Key Terms Glossary (2-column Table: Term | Definition)
    if pack.glossary:
        story.append(Paragraph("3. Key Terms & Concepts Glossary", h1_style))
        glossary_rows = [
            [Paragraph("Technical Term", table_header_style), Paragraph("Academic Definition", table_header_style)]
        ]
        for g in pack.glossary:
            term_p = Paragraph(f"<b>{html.escape(g.term)}</b>", table_cell_style)
            def_p = Paragraph(html.escape(g.definition), table_cell_style)
            glossary_rows.append([term_p, def_p])

        glossary_table = Table(glossary_rows, colWidths=[140, letter[0] - 72 - 140])
        glossary_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), PRIMARY),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
            ("BOTTOMPADDING", (0, 0), (-1, 0), 5),
            ("TOPPADDING", (0, 0), (-1, 0), 5),
            ("GRID", (0, 0), (-1, -1), 0.5, BOX_BORDER),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, BOX_BG]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ]))
        story.append(glossary_table)
        story.append(Spacer(1, 10))

    # 5. Practice Multiple-Choice Questions (20 MCQs)
    if pack.mcqs:
        story.append(Paragraph(f"4. Practice Multiple-Choice Examination ({len(pack.mcqs)} Questions)", h1_style))
        for idx, mcq in enumerate(pack.mcqs, start=1):
            q_elements = []
            safe_q = html.escape(mcq.question)
            q_elements.append(Paragraph(f"<b>Q{idx}.&nbsp; {safe_q}</b>", h2_style))

            for o_idx, opt in enumerate(mcq.options):
                letter_choice = chr(65 + o_idx)
                safe_opt = html.escape(str(opt).strip())
                prefix = f"<b>[{letter_choice}]</b>"
                q_elements.append(Paragraph(f"&nbsp;&nbsp;&nbsp;&nbsp;{prefix}&nbsp; {safe_opt}", bullet_style))

            # Answer and Explanation box
            ans_text = html.escape(mcq.correct_answer or (mcq.options[mcq.correct_index] if (mcq.correct_index is not None and 0 <= mcq.correct_index < len(mcq.options)) else "See explanation"))
            safe_expl = html.escape(mcq.explanation)
            expl_p = Paragraph(f"<b>Correct Answer:</b> {ans_text}<br/><b>Rationale:</b> {safe_expl}", table_cell_style)
            expl_table = Table([[expl_p]], colWidths=[letter[0] - 72])
            expl_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F0FDF4")),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#86EFAC")),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ]))
            q_elements.append(Spacer(1, 2))
            q_elements.append(expl_table)
            q_elements.append(Spacer(1, 6))
            story.append(KeepTogether(q_elements))

    # 6. Short-Answer Questions (5 Questions with Model Answers & Rubrics)
    if pack.short_answers:
        story.append(Paragraph(f"5. Short-Answer Conceptual Examination ({len(pack.short_answers)} Questions)", h1_style))
        for idx, sa in enumerate(pack.short_answers, start=1):
            sa_elements = []
            safe_q = html.escape(sa.question)
            sa_elements.append(Paragraph(f"<b>Q{idx}.&nbsp; {safe_q}</b>", h2_style))

            safe_ans = html.escape(sa.model_answer).replace("\n", "<br/>")
            ans_p = Paragraph(f"<b>Exemplar Model Answer:</b><br/>{safe_ans}", body_style)
            ans_table = Table([[ans_p]], colWidths=[letter[0] - 72])
            ans_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
                ("BOX", (0, 0), (-1, -1), 0.75, BOX_BORDER),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ]))
            sa_elements.append(ans_table)

            if sa.grading_criteria:
                sa_elements.append(Spacer(1, 4))
                sa_elements.append(Paragraph("<b>Grading Rubric & Key Concepts Required:</b>", h2_style))
                for crit in sa.grading_criteria:
                    safe_crit = html.escape(str(crit).strip())
                    sa_elements.append(Paragraph(f"&bull;&nbsp; {safe_crit}", bullet_style))

            sa_elements.append(Spacer(1, 8))
            story.append(KeepTogether(sa_elements))

    doc.build(story, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    return buffer


__all__ = [
    "export_notes_to_pdf",
    "export_notes_to_docx",
    "export_mcqs_to_csv",
    "export_pack_to_pdf",
    "NormalizedNotes",
    "NormalizedSection",
]
